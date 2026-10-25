import io
import re

import pytest
from openpyxl import load_workbook

from report_automator import InputFileError, load_config, run_report
from report_automator.cli import main

from .conftest import DEMO_CONFIG, HEADER, SAMPLES, d

SAMPLE_FILES = [
    SAMPLES / "jul_2026_sales_register.xlsx",
    SAMPLES / "aug_2026_sales.csv",
    SAMPLES / "sep_2026_sales.xlsx",
]


@pytest.fixture(scope="module")
def sample_report():
    return run_report(SAMPLE_FILES, load_config(DEMO_CONFIG))


def test_sample_data_totals_match_independent_calculation(sample_report):
    # Expected values were computed separately from the raw files with plain
    # pandas (see docs/case-study.md, "How the numbers were checked").
    s = sample_report.summary
    assert s.net_revenue == pytest.approx(518_322.00)
    assert list(s.by_month["revenue"]) == pytest.approx([170_257.0, 186_409.0, 161_656.0])
    assert s.bill_count == 1840
    c = sample_report.cleaned
    assert (c.rows_read, c.rows_kept) == (3612, 3566)


def test_every_removed_row_is_explained(sample_report):
    c = sample_report.cleaned
    per_row_removals = [i for i in c.issues if i.severity == "Removed" and i.row is not None]
    assert len(per_row_removals) == c.rows_removed


def test_excel_workbook_contents(sample_report):
    wb = load_workbook(io.BytesIO(sample_report.excel_bytes))
    assert wb.sheetnames == ["Summary", "By Product", "By Day", "By Week & Month",
                             "Payments & Categories", "Clean Data", "Issues"]
    summary = wb["Summary"]
    assert summary["A1"].value == "Namma Angadi Provisions (demo): Sales Report"
    assert summary["B6"].value == "₹5,18,322.00"
    assert len(summary._charts) == 2
    assert wb["Clean Data"].max_row == 3566 + 1
    assert wb["Issues"].max_row == len(sample_report.cleaned.issues) + 1
    assert wb["Clean Data"]["G2"].number_format.startswith("[>=10000000]")


def test_pdf_is_single_page(sample_report):
    pdf = sample_report.pdf_bytes
    assert pdf.startswith(b"%PDF")
    assert len(re.findall(rb"/Type\s*/Page\b", pdf)) == 1


def test_run_report_writes_files(tmp_path, write_xlsx, config):
    path = write_xlsx([HEADER, [d(1), "B1", "Milk", 1, 28, 28, "Cash"]])
    r = run_report([path], config, out_dir=tmp_path / "out")
    assert r.excel_path.exists() and r.pdf_path.exists()
    assert r.excel_path.name == "test-store_sales-report_2026-07-01_to_2026-07-01.xlsx"


def test_run_report_accepts_uploads(config):
    buf = io.BytesIO(b"Date,Item,Amount\n01/07/2026,Milk,28\n")
    r = run_report([("upload.csv", buf)], config)
    assert r.summary.net_revenue == 28


def test_same_file_twice_is_rejected(write_xlsx, config):
    path = write_xlsx([HEADER, [d(1), "B1", "Milk", 1, 28, 28, "Cash"]])
    with pytest.raises(InputFileError, match="given twice"):
        run_report([path, path], config)


def test_no_files(config):
    with pytest.raises(InputFileError, match="No input files"):
        run_report([], config)


def test_cli_success(tmp_path, capsys):
    code = main([*map(str, SAMPLE_FILES), "-c", str(DEMO_CONFIG), "-o", str(tmp_path)])
    out = capsys.readouterr().out
    assert code == 0
    assert "₹5,18,322.00" in out and "Reports written" in out
    assert len(list(tmp_path.glob("*.xlsx"))) == 1 and len(list(tmp_path.glob("*.pdf"))) == 1


def test_cli_quiet_prints_only_paths(tmp_path, capsys, write_xlsx):
    path = write_xlsx([HEADER, [d(1), "B1", "Milk", 1, 28, 28, "Cash"]])
    assert main([str(path), "-o", str(tmp_path), "-q"]) == 0
    lines = capsys.readouterr().out.strip().splitlines()
    assert len(lines) == 2 and all(line.strip().endswith((".xlsx", ".pdf")) for line in lines)


def test_cli_bad_input_exit_code(tmp_path, capsys):
    assert main([str(tmp_path / "missing.xlsx")]) == 2
    assert "not found" in capsys.readouterr().err


def test_cli_bad_config_exit_code(tmp_path, capsys, write_xlsx):
    cfg = tmp_path / "bad.yaml"
    cfg.write_text("colour: blue\n")
    path = write_xlsx([HEADER, [d(1), "B1", "Milk", 1, 28, 28, "Cash"]])
    assert main([str(path), "-c", str(cfg)]) == 2
    assert "Unknown setting" in capsys.readouterr().err


def test_cli_missing_config_file(tmp_path, capsys, write_xlsx):
    path = write_xlsx([HEADER, [d(1), "B1", "Milk", 1, 28, 28, "Cash"]])
    assert main([str(path), "-c", str(tmp_path / "nope.yaml")]) == 2


def test_cli_nothing_to_report(tmp_path, capsys, write_xlsx):
    path = write_xlsx([["Date", "Item", "Amount"], ["garbage", "Milk", 28]])
    assert main([str(path), "-o", str(tmp_path)]) == 1
    assert "Nothing to report" in capsys.readouterr().err
