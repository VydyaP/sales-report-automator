"""Formatted Excel workbook: summary with charts, breakdowns, clean data and
the issue log. Built with openpyxl so charts are native Excel charts the
owner can edit, not pasted images."""

from __future__ import annotations

import datetime as dt
from pathlib import Path
from typing import BinaryIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .cleaner import CHECK, FIXED, REMOVED, CleanResult
from .summarize import Summary

BRAND = "1F4E5F"
LIGHT = "E8F1F3"
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=BRAND)
TITLE_FONT = Font(bold=True, size=16, color=BRAND)
SUBTLE_FONT = Font(italic=True, color="666666")
THIN = Side(style="thin", color="D0D7DA")
SEVERITY_FILL = {
    REMOVED: PatternFill("solid", fgColor="FDE2E1"),
    CHECK: PatternFill("solid", fgColor="FFF4CE"),
    FIXED: PatternFill("solid", fgColor="E3F4E8"),
}


def money_format(symbol: str) -> str:
    """Indian lakh/crore grouping (12,34,567.00) in Excel's own format language."""
    s = f'"{symbol}"'
    return (
        f'[>=10000000]{s}##\\,##\\,##\\,##0.00;'
        f'[>=100000]{s}##\\,##\\,##0.00;'
        f'{s}##,##0.00'
    )


DATE_FMT = "dd-mmm-yyyy"
PCT_FMT = "0.0%"
QTY_FMT = "#,##0.##"


def build_workbook(result: CleanResult, summary: Summary) -> Workbook:
    money = money_format(summary.currency)
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"

    # Sheets that charts depend on are written first.
    product_ws = wb.create_sheet("By Product")
    prod_cols = [("product", "Product", None, 34)]
    if "category" in summary.by_product:
        prod_cols.append(("category", "Category", None, 16))
    if "units" in summary.by_product:
        prod_cols.append(("units", "Units sold", QTY_FMT, 12))
    prod_cols += [("revenue", "Sales", money, 16), ("share", "Share of sales", PCT_FMT, 14),
                  ("lines", "Sale lines", "0", 11)]
    _write_table(product_ws, summary.by_product, prod_cols, title="Sales by product")

    day_ws = wb.create_sheet("By Day")
    day_cols = [("date", "Date", DATE_FMT, 14), ("weekday", "Weekday", None, 12),
                ("revenue", "Sales", money, 16)]
    if "bills" in summary.by_day:
        day_cols.append(("bills", "Bills", "0", 9))
    day_cols.append(("lines", "Sale lines", "0", 11))
    _write_table(day_ws, summary.by_day, day_cols, title="Sales by day")

    period_ws = wb.create_sheet("By Week & Month")
    end_row = _write_table(period_ws, summary.by_month,
                           [("month", "Month", "mmm-yyyy", 16), ("revenue", "Sales", money, 16),
                            ("lines", "Sale lines", "0", 11)], title="Sales by month")
    _write_table(period_ws, summary.by_week,
                 [("week_starting", "Week starting (Mon)", DATE_FMT, 20), ("revenue", "Sales", money, 16),
                  ("lines", "Sale lines", "0", 11)], title="Sales by week", start_row=end_row + 2)
    _write_table(period_ws, summary.by_weekday,
                 [("weekday", "Weekday", None, 12), ("avg_revenue", "Average daily sales", money, 20),
                  ("days", "Days in period", "0", 14)],
                 title="Average by weekday", start_row=1, start_col=5)

    if summary.by_payment_mode is not None or summary.by_category is not None:
        mix_ws = wb.create_sheet("Payments & Categories")
        row = 1
        if summary.by_payment_mode is not None:
            row = _write_table(mix_ws, summary.by_payment_mode,
                               [("payment_mode", "Payment method", None, 20), ("revenue", "Sales", money, 16),
                                ("share", "Share", PCT_FMT, 10), ("lines", "Sale lines", "0", 11)],
                               title="Sales by payment method") + 2
        if summary.by_category is not None:
            _write_table(mix_ws, summary.by_category,
                         [("category", "Category", None, 20), ("revenue", "Sales", money, 16),
                          ("share", "Share", PCT_FMT, 10), ("lines", "Sale lines", "0", 11)],
                         title="Sales by category", start_row=row)

    data_ws = wb.create_sheet("Clean Data")
    _write_table(data_ws, result.data, [
        ("date", "Date", DATE_FMT, 13), ("invoice_id", "Invoice no", None, 12),
        ("product", "Product", None, 32), ("category", "Category", None, 14),
        ("quantity", "Qty", QTY_FMT, 8), ("unit_price", "Rate", money, 12),
        ("amount", "Amount", money, 14), ("payment_mode", "Payment", None, 12),
        ("source_file", "Source file", None, 26), ("source_row", "Source row", "0", 11),
    ], title=None, autofilter=True)

    issues_ws = wb.create_sheet("Issues")
    issues = result.issues_frame()
    _write_table(issues_ws, issues, [
        ("severity", "Type", None, 10), ("file", "File", None, 26), ("row", "Row", "0", 7),
        ("field", "Field", None, 13), ("message", "What happened", None, 80),
        ("original_value", "Original value", None, 22),
    ], title=None, autofilter=True)
    for r in range(2, issues_ws.max_row + 1):
        fill = SEVERITY_FILL.get(issues_ws.cell(r, 1).value)
        if fill:
            issues_ws.cell(r, 1).fill = fill
        issues_ws.cell(r, 5).alignment = Alignment(wrap_text=True, vertical="top")

    _write_summary_sheet(ws, result, summary, product_ws, day_ws, money)
    return wb


def save_workbook(wb: Workbook, target: str | Path | BinaryIO) -> None:
    wb.save(target)


def _write_summary_sheet(ws: Worksheet, result: CleanResult, s: Summary,
                         product_ws: Worksheet, day_ws: Worksheet, money: str) -> None:
    ws["A1"] = f"{s.business_name}: Sales Report"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = s.period_label
    ws["A2"].font = Font(size=12, color="333333")
    ws["A3"] = f"Generated {dt.date.today().strftime('%d %b %Y')} by Sales Report Automator"
    ws["A3"].font = SUBTLE_FONT

    row = 5
    _section(ws, row, "Key numbers")
    row += 1
    for label, value in s.kpis():
        ws.cell(row, 1, label).font = Font(bold=True)
        ws.cell(row, 2, value).alignment = Alignment(horizontal="right")
        for c in (1, 2):
            ws.cell(row, c).border = Border(bottom=THIN)
        row += 1

    row += 1
    _section(ws, row, "What stands out")
    row += 1
    for line in s.insights:
        cell = ws.cell(row, 1, f"• {line}")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
        ws.row_dimensions[row].height = 32
        row += 1

    row += 1
    _section(ws, row, "Data quality")
    row += 1
    quality = [
        ("Rows read", result.rows_read),
        ("Rows used in this report", result.rows_kept),
        ("Rows removed", result.rows_removed),
        ("Automatic corrections", result.count(FIXED)),
        ("Items to double-check", result.count(CHECK)),
    ]
    for label, value in quality:
        ws.cell(row, 1, label)
        ws.cell(row, 2, value).alignment = Alignment(horizontal="right")
        row += 1
    ws.cell(row, 1, "See the 'Issues' sheet for every change, row by row.").font = SUBTLE_FONT

    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 34

    # Charts reference the breakdown sheets (header on row 2, data from row 3).
    n_products = min(len(s.by_product), 10)
    rev_col = _col_index(product_ws, "Sales")
    bar = BarChart()
    bar.type = "bar"
    bar.title = f"Top {n_products} products by sales"
    bar.y_axis.title = None
    bar.legend = None
    bar.add_data(Reference(product_ws, min_col=rev_col, min_row=2, max_row=2 + n_products), titles_from_data=True)
    bar.set_categories(Reference(product_ws, min_col=1, min_row=3, max_row=2 + n_products))
    bar.x_axis.scaling.orientation = "maxMin"  # biggest at the top
    bar.y_axis.numFmt = '"' + s.currency + '"#,##0'
    bar.height, bar.width = 9, 16
    ws.add_chart(bar, "D5")

    day_rev_col = _col_index(day_ws, "Sales")
    line = LineChart()
    line.title = "Daily sales"
    line.legend = None
    line.add_data(Reference(day_ws, min_col=day_rev_col, min_row=2, max_row=2 + len(s.by_day)), titles_from_data=True)
    line.set_categories(Reference(day_ws, min_col=1, min_row=3, max_row=2 + len(s.by_day)))
    line.series[0].smooth = False  # straight segments: don't invent values between days
    line.x_axis.number_format = "dd-mmm"
    line.y_axis.numFmt = '"' + s.currency + '"#,##0'
    line.height, line.width = 8, 16
    ws.add_chart(line, "D24")


def _section(ws: Worksheet, row: int, text: str) -> None:
    cell = ws.cell(row, 1, text)
    cell.font = Font(bold=True, size=12, color=BRAND)
    cell.fill = PatternFill("solid", fgColor=LIGHT)
    ws.cell(row, 2).fill = PatternFill("solid", fgColor=LIGHT)


def _write_table(ws: Worksheet, df: pd.DataFrame, columns: list[tuple[str, str, str | None, int]],
                 title: str | None, start_row: int = 1, start_col: int = 1,
                 autofilter: bool = False) -> int:
    """Write df[columns] as a styled table. Returns the last row written."""
    row = start_row
    if title:
        ws.cell(row, start_col, title).font = Font(bold=True, size=12, color=BRAND)
        row += 1
    header_row = row
    for j, (_, label, _, width) in enumerate(columns):
        c = ws.cell(row, start_col + j, label)
        c.font, c.fill = HEADER_FONT, HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        letter = get_column_letter(start_col + j)
        ws.column_dimensions[letter].width = max(ws.column_dimensions[letter].width or 0, width)

    for record in df.to_dict("records"):
        row += 1
        for j, (key, _, fmt, _) in enumerate(columns):
            value = record.get(key)
            value = _excel_value(value)
            c = ws.cell(row, start_col + j, value)
            if fmt and value is not None:
                c.number_format = fmt
            c.border = Border(bottom=THIN)

    if title is None:
        ws.freeze_panes = ws.cell(header_row + 1, start_col)
    if autofilter and row > header_row:
        ws.auto_filter.ref = (f"{get_column_letter(start_col)}{header_row}:"
                              f"{get_column_letter(start_col + len(columns) - 1)}{row}")
    return row


def _excel_value(value):
    if value is None or value is pd.NA or value is pd.NaT:
        return None
    if isinstance(value, float) and pd.isna(value):
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if hasattr(value, "item"):  # numpy scalar
        return value.item()
    return value


def _col_index(ws: Worksheet, header: str) -> int:
    for cell in ws[2]:
        if cell.value == header:
            return cell.column
    raise KeyError(header)  # pragma: no cover - programming error
