"""Build the images used in the README from the real sample files.

    python scripts/make_doc_images.py

* docs/images/before.png      : real problem rows from the sample exports, labelled
* docs/images/report_pdf.png  : page 1 of the generated PDF   (needs `pdftoppm`, from poppler)
* docs/images/report_excel.png: Excel summary + charts        (needs LibreOffice `soffice` and `pdftoppm`)
"""

from __future__ import annotations

import csv
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from report_automator import load_config, run_report  # noqa: E402

SAMPLES = ROOT / "sample_data"
IMAGES = ROOT / "docs" / "images"

# (file, row number as the owner sees it, what is wrong)
EXAMPLES = [
    ("jul_2026_sales_register.xlsx", 1, "Title rows above the real header"),
    ("jul_2026_sales_register.xlsx", 45, "Impossible date (31 June)"),
    ("jul_2026_sales_register.xlsx", 53, "'Day Total' line: would double-count"),
    ("jul_2026_sales_register.xlsx", 314, "Same line keyed in twice"),
    ("aug_2026_sales.csv", 2, "Re-export of a July bill (overlap)"),
    ("aug_2026_sales.csv", 62, "Amount missing"),
    ("aug_2026_sales.csv", 152, "Date missing"),
    ("aug_2026_sales.csv", 262, "Amount typed as 'N/A'"),
    ("aug_2026_sales.csv", 263, "Qty and rate missing"),
    ("sep_2026_sales.xlsx", 30, "Month 13 does not exist"),
    ("sep_2026_sales.xlsx", 1134, "Grand Total row"),
]


def read_row(name: str, row: int) -> list[str]:
    path = SAMPLES / name
    if path.suffix == ".csv":
        with open(path, encoding="utf-8") as f:
            values = list(csv.reader(f))[row - 1]
    else:
        ws = load_workbook(path, read_only=True).active
        values = [c for c in next(ws.iter_rows(min_row=row, max_row=row, values_only=True))]
    out = []
    for v in values:
        if v is None:
            out.append("")
        elif hasattr(v, "strftime"):
            out.append(v.strftime("%d-%m-%Y"))
        else:
            out.append(str(v))
    return out


def make_before() -> None:
    header = ["File : row", "Date", "Product", "Amount", "Payment", "Problem"]
    rows = []
    for name, row, problem in EXAMPLES:
        v = read_row(name, row) + [""] * 9
        short = {"jul_2026_sales_register.xlsx": "jul.xlsx", "aug_2026_sales.csv": "aug.csv",
                 "sep_2026_sales.xlsx": "sep.xlsx"}[name]
        if name.startswith("sep"):
            date, product, amount, pay = v[0], v[2], v[7], v[8]
        else:
            date, product, amount, pay = v[0], v[2], v[6], v[7]
        if row == 1 and name.startswith("jul"):
            date, product, amount, pay = v[0][:20] + "…", "", "", ""
        rows.append([f"{short} : {row}", date, product, amount, pay, problem])

    fig, ax = plt.subplots(figsize=(13, 0.33 * (len(rows) + 1)), dpi=150)
    ax.axis("off")
    table = ax.table(cellText=rows, colLabels=header, loc="upper center", cellLoc="left",
                     colWidths=[0.12, 0.17, 0.19, 0.09, 0.08, 0.31])
    table.auto_set_font_size(False)
    table.set_fontsize(9.5)
    table.scale(1, 1.45)
    for (r, c), cell in table.get_celld().items():
        cell.set_edgecolor("#D0D7DA")
        if r == 0:
            cell.set_facecolor("#1F4E5F")
            cell.set_text_props(color="white", weight="bold")
        elif c == 5:
            cell.set_facecolor("#FDE2E1")
    ax.set_title("Before: real rows from the three sample exports", loc="left", fontsize=12,
                 fontweight="bold", color="#1F4E5F")
    fig.tight_layout()
    fig.savefig(IMAGES / "before.png", facecolor="white", bbox_inches="tight", pad_inches=0.15)
    plt.close(fig)


def make_report_previews() -> None:
    if not shutil.which("pdftoppm"):
        print("pdftoppm not found: skipping report previews")
        return
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        files = sorted(SAMPLES.glob("*.xlsx")) + sorted(SAMPLES.glob("*.csv"))
        result = run_report(files, load_config(ROOT / "configs" / "demo_store.yaml"), out_dir=tmp)
        subprocess.run(["pdftoppm", "-png", "-r", "110", "-singlefile", str(result.pdf_path),
                        str(IMAGES / "report_pdf")], check=True)
        if shutil.which("soffice"):
            subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(tmp / "xl"),
                            str(result.excel_path)], check=True, capture_output=True, timeout=180)
            xl_pdf = next((tmp / "xl").glob("*.pdf"))
            subprocess.run(["pdftoppm", "-png", "-r", "80", "-f", "1", "-l", "2", str(xl_pdf), str(tmp / "xl" / "p")],
                           check=True)
            pages = sorted((tmp / "xl").glob("p*.png"))
            imgs = [plt.imread(p) for p in pages[:2]]
            fig, axes = plt.subplots(1, len(imgs), figsize=(6 * len(imgs), 8), dpi=110)
            for ax, img in zip(axes if len(imgs) > 1 else [axes], imgs):
                ax.imshow(img)
                ax.axis("off")
            fig.tight_layout()
            fig.savefig(IMAGES / "report_excel.png", facecolor="white")
            plt.close(fig)
        else:
            print("soffice not found: skipping Excel preview")


if __name__ == "__main__":
    IMAGES.mkdir(parents=True, exist_ok=True)
    make_before()
    make_report_previews()
    for p in sorted(IMAGES.glob("*.png")):
        print("wrote", p.relative_to(ROOT))
